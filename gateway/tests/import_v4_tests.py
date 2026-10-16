"""Knowledge import pipeline tests (contract §1) — deterministic, offline.

Run:  gateway/.venv-v4/bin/python gateway/tests/import_v4_tests.py
"""

import os
import sys
import tempfile
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from modules._jsonstore import ConfirmRequired, StoreUnavailable  # noqa: E402
from modules.knowledge import parsers  # noqa: E402
from modules.knowledge.imports import (ImportStore, JobExpired,  # noqa: E402
                                       JobNotCancellable, JobNotRetryable)
from modules.knowledge.store import (DuplicateContent, KnowledgeStore,  # noqa: E402
                                     MAX_DOC_CHARS)

PASS, FAIL = [], []


def check(name, cond, extra=""):
    (PASS if cond else FAIL).append(name)
    print(("PASS " if cond else "FAIL ") + name + (("  → " + str(extra)) if extra and not cond else ""))


MIN_PDF = b"""%PDF-1.4
1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj
2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj
3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 300 200]/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>endobj
4 0 obj<</Length 52>>stream
BT /F1 14 Tf 20 100 Td (Hello Knowledge Base) Tj ET
endstream
endobj
5 0 obj<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>endobj
trailer<</Root 1 0 R>>
%%EOF
"""


def make_docx(path, text):
    import docx
    d = docx.Document()
    d.add_paragraph(text)
    d.save(path)


def make_pptx(path, text):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    tb.text_frame.text = text
    prs.save(path)


def main():
    tmp = tempfile.mkdtemp(prefix="v4imp_")
    kb_path = os.path.join(tmp, "kb.json")
    imp_path = os.path.join(tmp, "imports.json")
    kstore = KnowledgeStore(kb_path).load()
    ds = kstore.create_dataset("课件 · 测试")
    isr = ImportStore(imp_path, scratch_dir=os.path.join(tmp, "scratch")).load()

    # ---- 1. upload ≠ 入库 ------------------------------------------------
    txt = "check-in, check-out, reservation, front desk, luggage"
    job = isr.create_job(dataset_id=ds["id"], filename="unit3.txt",
                         data=txt.encode("utf-8"), title="Unit 3 酒店")
    check("txt 上传后为 preview_ready", job["status"] == "preview_ready", job.get("status"))
    check("上传成功但知识库仍为空（上传≠入库）",
          len(KnowledgeStore(kb_path).load().list_documents()) == 0)
    check("预览内容可用", job["preview"].startswith("check-in"))
    check("预览标记未截断", job["preview_truncated"] is False)
    check("记录 content_hash", bool(job["content_hash"]))

    # ---- 2. commit → applied + read-back --------------------------------
    try:
        isr.commit(job["job_id"], knowledge_store=kstore, confirm=False)
        check("commit 需 confirm", False, "未拒绝")
    except ConfirmRequired:
        check("commit 需 confirm", True)
    res = isr.commit(job["job_id"], knowledge_store=kstore, confirm=True)
    check("commit 后为 applied", res["job"]["status"] == "applied")
    check("commit 报告 readback_ok", res["readback_ok"] is True)
    docs = KnowledgeStore(kb_path).load().list_documents(dataset_id=ds["id"])
    check("文档已真实入库且可读回", len(docs) == 1 and docs[0]["chunk_count"] >= 1)
    check("source_type 记录为 text", docs[0]["source"]["type"] == "text")

    # ---- 3. duplicate content -------------------------------------------
    dup = isr.create_job(dataset_id=ds["id"], filename="unit3-copy.txt",
                         data=txt.encode("utf-8"))
    try:
        isr.commit(dup["job_id"], knowledge_store=kstore, confirm=True)
        check("重复内容被拒绝", False, "未拒绝")
    except DuplicateContent:
        check("重复内容被拒绝", True)
    check("重复后 job 仍为 preview_ready（可换资料集）",
          isr.get_job(dup["job_id"])["status"] == "preview_ready")
    check("重复未写入第二条文档",
          len(KnowledgeStore(kb_path).load().list_documents()) == 1)

    # ---- 4. cancel ------------------------------------------------------
    c = isr.cancel(dup["job_id"], confirm=True)
    check("preview_ready 可取消", c["status"] == "cancelled")
    try:
        isr.cancel(job["job_id"], confirm=True)
        check("applied 不可取消", False, "未拒绝")
    except JobNotCancellable:
        check("applied 不可取消", True)

    # ---- 5. unsupported / non-P0 格式 -----------------------------------
    bad = isr.create_job(dataset_id=ds["id"], filename="image.png",
                         data=b"\x89PNG\r\n")
    check("非 P0 格式 → unsupported_content_type",
          bad["status"] == "failed" and bad["error_code"] == "unsupported_content_type",
          bad.get("error_code"))
    legacy = isr.create_job(dataset_id=ds["id"], filename="old.doc", data=b"\xd0\xcf")
    check(".doc 归为 unsupported_content_type",
          legacy["error_code"] == "unsupported_content_type")

    # ---- 6. 超长文档：拒绝而非截断 --------------------------------------
    long_job = isr.create_job(dataset_id=ds["id"], filename="huge.txt",
                              data=("长" * (MAX_DOC_CHARS + 10)).encode("utf-8"))
    check("超长文档 → cap_exceeded（拒绝）",
          long_job["status"] == "failed" and long_job["error_code"] == "cap_exceeded",
          long_job.get("error_code"))
    check("超长文档未产生任何入库",
          len(KnowledgeStore(kb_path).load().list_documents()) == 1)

    # ---- 7. md / docx / pptx / pdf 真实解析 ------------------------------
    md = isr.create_job(dataset_id=ds["id"], filename="travel.md",
                        data=b"# Travel\nWhere is the station?")
    check("md 解析成功", md["status"] == "preview_ready", md.get("error_code"))
    dx = os.path.join(tmp, "worksheet.docx")
    make_docx(dx, "Reservation and front desk vocabulary")
    docx_job = isr.create_job(dataset_id=ds["id"], filename="worksheet.docx",
                              data=Path(dx).read_bytes())
    check("docx 解析成功", docx_job["status"] == "preview_ready", docx_job.get("error_code"))
    check("docx 抽取到文本",
          "Reservation" in (docx_job.get("preview") or ""), docx_job.get("preview"))
    px = os.path.join(tmp, "slides.pptx")
    make_pptx(px, "Hello from slide")
    pptx_job = isr.create_job(dataset_id=ds["id"], filename="slides.pptx",
                              data=Path(px).read_bytes())
    check("pptx 解析成功", pptx_job["status"] == "preview_ready", pptx_job.get("error_code"))
    check("pptx 抽取到文本",
          "Hello from slide" in (pptx_job.get("preview") or ""))
    pdf_job = isr.create_job(dataset_id=ds["id"], filename="textbook.pdf", data=MIN_PDF)
    check("pdf 解析成功", pdf_job["status"] == "preview_ready", pdf_job.get("error_code"))
    check("pdf 抽取到文本", "Hello Knowledge Base" in (pdf_job.get("preview") or ""))
    check("pdf source_type=pdf_text", pdf_job["source_type"] == "pdf_text")
    kinds = sorted({c["ext"]: c["available"] for c in parsers.capabilities()}.items())
    check("六种解析器均可用（五种 P0 + .epub）",
          all(v for _, v in kinds) and len(kinds) == 6, kinds)

    # commit the docx one to prove docx round-trips into the store
    r2 = isr.commit(docx_job["job_id"], knowledge_store=kstore, confirm=True)
    check("docx commit 后 applied", r2["job"]["status"] == "applied")
    d2 = KnowledgeStore(kb_path).load().get_document(r2["document"]["id"])
    check("docx 文档 source_type=docx_text", d2["source"]["type"] == "docx_text")

    # ---- 8. 过期 ---------------------------------------------------------
    isr.ttl = 0
    short = isr.create_job(dataset_id=ds["id"], filename="short.txt",
                           data=b"no duplicates here at all xyz")
    time.sleep(0.01)
    try:
        isr.commit(short["job_id"], knowledge_store=kstore, confirm=True)
        check("超期预览 → job_expired", False, "未拒绝")
    except JobExpired:
        check("超期预览 → job_expired", True)
    isr.ttl = 1800

    # ---- 9. retry 路径 ---------------------------------------------------
    broken = isr.create_job(dataset_id=ds["id"], filename="broken.pdf",
                            data=b"not really a pdf")
    check("损坏 pdf → extract_failed（可重试）",
          broken["status"] == "failed" and broken["error_code"] == "extract_failed",
          broken.get("error_code"))
    again = isr.retry(broken["job_id"], confirm=True)
    check("retry 重新解析并再次失败", again["status"] == "failed")
    try:
        isr.retry(job["job_id"], confirm=True)
        check("applied 不可 retry", False, "未拒绝")
    except JobNotRetryable:
        check("applied 不可 retry", True)

    # ---- 10. 列表与幂等 --------------------------------------------------
    lst = isr.list_jobs(status="applied")
    check("按状态过滤任务列表", len(lst) == 2, len(lst))
    idem = isr.create_job(dataset_id=ds["id"], filename="idem.txt",
                          data=b"idempotent content", idempotency_key="k-1")
    idem2 = isr.create_job(dataset_id=ds["id"], filename="idem.txt",
                           data=b"idempotent content", idempotency_key="k-1")
    check("相同 Idempotency-Key 返回同一 job", idem["job_id"] == idem2["job_id"])

    # ---- 11. import store 版本门禁 --------------------------------------
    Path(imp_path).write_text('{"schema_version": 99, "jobs": {}}', "utf-8")
    un = ImportStore(imp_path).load()
    check("导入 store 高版本拒写", not un.available)
    original = Path(imp_path).read_text("utf-8")
    try:
        un._persist()
        check("导入 store unavailable 拒写", False, "未拒绝")
    except StoreUnavailable:
        check("导入 store unavailable 拒写", True)
    check("导入 store 原文件未被覆盖", Path(imp_path).read_text("utf-8") == original)

    print("\n==== SUMMARY ====")
    print("total=%d pass=%d fail=%d" % (len(PASS) + len(FAIL), len(PASS), len(FAIL)))
    if FAIL:
        print("FAILED: " + ", ".join(FAIL))
    return 1 if FAIL else 0


if __name__ == "__main__":
    sys.exit(main())
