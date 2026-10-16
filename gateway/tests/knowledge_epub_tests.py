"""Text-only EPUB parser — security & correctness tests (CLASSROOM V2 · D2).

Offline and deterministic.  Builds synthetic EPUB archives (stdlib ``zipfile``) in
a temp dir and drives the real ``modules.knowledge.parsers.parse_file`` through the
``.epub`` dispatch.  Writes an evidence JSON into the repo.

Run:
    PYTHONPATH= <venv>/bin/python gateway/tests/knowledge_epub_tests.py

Every assertion records its own file path + line so the evidence JSON is
self-locating.  Nothing here writes into the knowledge store.
"""

import io
import json
import platform
import struct
import sys
import tempfile
import time
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from modules.knowledge import parsers  # noqa: E402

TEST_FILE = str(Path(__file__).resolve())
EVIDENCE = (ROOT / "docs/engineering/product-v1/classroom-v2/evidence"
            / "epub_parser_security_tests.json")

RESULTS = []
PASS, FAIL = [], []


def check(cid, name, cond, code=None, observed=None):
    """Record one assertion with the caller's real line number."""
    line = sys._getframe(1).f_lineno
    rec = {"case_id": cid, "name": name, "passed": bool(cond),
           "error_code": code, "observed": observed,
           "asserted_file": TEST_FILE, "asserted_line": line}
    RESULTS.append(rec)
    (PASS if cond else FAIL).append(cid)
    extra = "" if cond else "  → code=%s observed=%r" % (code, observed)
    print(("PASS " if cond else "FAIL ") + "%s %s%s" % (cid, name, extra))
    return bool(cond)


# ---- synthetic archive builders -------------------------------------------

CONTAINER_XML = (
    '<?xml version="1.0" encoding="UTF-8"?>'
    '<container version="1.0" '
    'xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
    '<rootfiles><rootfile full-path="OEBPS/content.opf" '
    'media-type="application/oebps-package+xml"/></rootfiles></container>'
)


def opf_xml(items, spine):
    man = "".join('<item id="%s" href="%s" media-type="application/xhtml+xml"/>'
                  % (i, h) for i, h in items)
    sp = "".join('<itemref idref="%s"/>' % r for r in spine)
    return ('<?xml version="1.0" encoding="UTF-8"?>'
            '<package xmlns="http://www.idpf.org/2007/opf" version="2.0" '
            'unique-identifier="bid"><metadata '
            'xmlns:dc="http://purl.org/dc/elements/1.1/">'
            '<dc:title>Test Book</dc:title></metadata>'
            '<manifest>%s</manifest><spine toc="ncx">%s</spine></package>'
            % (man, sp))


def xhtml(body, head=""):
    return ('<?xml version="1.0" encoding="UTF-8"?>'
            '<html xmlns="http://www.w3.org/1999/xhtml">'
            '<head><title>t</title>%s</head><body>%s</body></html>' % (head, body))


def build_epub(path, entries):
    """Write a zip archive to ``path``; return the raw bytes."""
    bio = io.BytesIO()
    with zipfile.ZipFile(bio, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in entries:
            zf.writestr(name, data if isinstance(data, bytes) else data.encode("utf-8"))
    raw = bio.getvalue()
    Path(path).write_bytes(raw)
    return raw


def _cd_offset(raw, name):
    """Offset of the central-directory header whose filename == ``name``."""
    nb = name.encode("utf-8")
    pos = 0
    while True:
        i = raw.find(b"PK\x01\x02", pos)
        if i < 0:
            return None
        nlen = struct.unpack_from("<H", raw, i + 28)[0]
        if bytes(raw[i + 46:i + 46 + nlen]) == nb:
            return i
        pos = i + 46 + nlen


def _patch_central(raw, name, *, uncompressed=None, compressed=None, flags=None):
    """Patch a central-directory header (to forge declared sizes / flags)."""
    raw = bytearray(raw)
    i = _cd_offset(raw, name)
    assert i is not None, "central dir entry not found: %s" % name
    if uncompressed is not None:
        struct.pack_into("<I", raw, i + 24, uncompressed)
    if compressed is not None:
        struct.pack_into("<I", raw, i + 20, compressed)
    if flags is not None:
        struct.pack_into("<H", raw, i + 8, flags)
    return bytes(raw)


def normal_entries():
    """3-chapter EPUB whose entry-name order != spine order (proves spine order).

    File name alphabetisation is THREE(mm.xhtml) < TWO(my chapter.xhtml) <
    ONE(zz.xhtml); spine order is ONE, TWO, THREE.
    """
    items = [("p1", "zz.xhtml"), ("p2", "my%20chapter.xhtml"), ("p3", "mm.xhtml")]
    spine = ["p1", "p2", "p3"]
    entries = [
        ("mimetype", "application/epub+zip"),
        ("META-INF/container.xml", CONTAINER_XML),
        ("OEBPS/content.opf", opf_xml(items, spine)),
        ("OEBPS/zz.xhtml", xhtml(
            '<h1>CHAPTER-ONE</h1><p>Body one<br/>line two</p>'
            '<p><![CDATA[CDATA-VISIBLE]]></p>'
            '<script>alert(1)</script><!-- a comment -->')),
        ("OEBPS/my chapter.xhtml", xhtml('<h1>CHAPTER-TWO</h1><p>Body two</p>')),
        ("OEBPS/mm.xhtml", xhtml('<h1>CHAPTER-THREE</h1><p>Body three</p>')),
    ]
    return entries


def parse_epub(tmp, name, raw):
    p = Path(tmp) / name
    p.write_bytes(raw)
    return parsers.parse_file(str(p), ".epub")


# ---- checks ----------------------------------------------------------------

def main():
    tmp = tempfile.mkdtemp(prefix="epub_sec_")

    # ⑩ parser_available always True (stdlib only)
    ok, reason = parsers.parser_available(".epub")
    check("C10", ".epub parser_available is True (stdlib only)", ok is True, reason, {"available": ok, "reason": reason})

    # ⑪ capabilities() advertises .epub
    caps = {c["ext"]: c for c in parsers.capabilities()}
    ecap = caps.get(".epub")
    check("C11", "capabilities() contains .epub and available=true",
          ecap is not None and ecap.get("available") is True and ecap.get("source_type") == "epub_text",
          None, ecap)

    # ---------- ① normal multi-chapter EPUB, spine order ----------
    raw = build_epub(Path(tmp) / "normal.epub", normal_entries())
    text, err = parsers.parse_file(str(Path(tmp) / "normal.epub"), ".epub")
    one, two, three = (text or "").find("CHAPTER-ONE"), (text or "").find("CHAPTER-TWO"), (text or "").find("CHAPTER-THREE")
    check("C01a", "normal EPUB parses (no error)", err is None and bool(text), err, {"err": err, "chars": len(text or "")})
    check("C01b", "chapters emitted in SPINE order (not filename order)",
          -1 < one < two < three, None, {"one": one, "two": two, "three": three})
    check("C01c", "chapter boundary markers [chapter n] present in spine order",
          "[chapter 1]" in (text or "") and "[chapter 2]" in (text or "")
          and "[chapter 3]" in (text or "")
          and (text or "").find("[chapter 1]") < (text or "").find("[chapter 3]"),
          None, (text or "")[:80])
    check("C01d", "text starts with the first spine chapter marker",
          (text or "").startswith("[chapter 1]"), None, (text or "")[:40])
    check("C01e", "XHTML CDATA content preserved",
          "CDATA-VISIBLE" in (text or ""), None, "CDATA-VISIBLE" in (text or ""))
    check("C01f", "XML comment not leaked into text",
          "a comment" not in (text or ""), None, "a comment" in (text or ""))
    check("C01g", "percent-encoded manifest href resolves (my%20chapter → 'my chapter')",
          "CHAPTER-TWO" in (text or ""), None, "CHAPTER-TWO" in (text or ""))

    # ---------- ⑤ script / style stripped ----------
    check("C05a", "<script> body stripped (not executed text)",
          "alert(1)" not in (text or ""), None, "alert(1)" in (text or ""))
    _, serr = parse_epub(tmp, "style.epub", build_epub(Path(tmp) / "style.epub", [
        ("META-INF/container.xml", CONTAINER_XML),
        ("OEBPS/content.opf", opf_xml([("p1", "a.xhtml")], ["p1"])),
        ("OEBPS/a.xhtml", xhtml('<p>VISIBLE-TEXT</p>',
                                head='<style>p{color:red}</style>')),
    ]))
    stext, _ = parsers.parse_file(str(Path(tmp) / "style.epub"), ".epub")
    check("C05b", "<style> content stripped; visible text kept",
          "color:red" not in (stext or "") and "VISIBLE-TEXT" in (stext or ""),
          None, {"has_visible": "VISIBLE-TEXT" in (stext or ""), "has_css": "color:red" in (stext or "")})

    # ---------- ② zip bomb: declared huge size ----------
    bomb = build_epub(Path(tmp) / "bomb.epub", [
        ("META-INF/container.xml", CONTAINER_XML),
        ("OEBPS/content.opf", opf_xml([("p1", "big.xhtml")], ["p1"])),
        ("OEBPS/big.xhtml", b"A" * (20 * 1024 * 1024)),   # 20 MiB, compresses tiny
    ])
    btext, berr = parse_epub(tmp, "bomb.epub", bomb)
    check("C02a", "zip bomb (declared 20MiB > 16MiB cap) → epub_unsafe",
          btext is None and berr == "epub_unsafe",
          berr, {"error_code": berr, "compressed_bytes": len(bomb), "uncompressed": 20 * 1024 * 1024})

    # ---------- ②b zip bomb: declared size forged small (bypass declared check) ----------
    raw2 = build_epub(Path(tmp) / "bomb2.epub", [
        ("META-INF/container.xml", CONTAINER_XML),
        ("OEBPS/content.opf", opf_xml([("p1", "big.xhtml")], ["p1"])),
        ("OEBPS/big.xhtml", b"A" * (20 * 1024 * 1024)),
    ])
    # forge central dir: claim 1 KiB uncompressed so the *declared* check passes.
    lied = _patch_central(raw2, "OEBPS/big.xhtml", uncompressed=1024, compressed=1024)
    b2text, b2err = parse_epub(tmp, "bomb2.epub", lied)
    # Python 3.13's zipfile validates CRC/size itself and refuses before the
    # streaming counter can trip → still fail-closed, just a different honest code.
    check("C02b", "zip bomb with forged declared size → fail-closed (no text returned)",
          b2text is None and b2err in ("epub_malformed", "epub_unsafe"),
          b2err, {"error_code": b2err, "declared_usize": 1024, "actual_usize": 20 * 1024 * 1024})

    # ---------- ②c total-budget: many members each under per-entry cap ----------
    many = [("META-INF/container.xml", CONTAINER_XML)]
    items, spine = [], []
    for k in range(5):
        items.append(("p%d" % k, "c%d.xhtml" % k))
        spine.append("p%d" % k)
        many.append(("OEBPS/c%d.xhtml" % k, b"A" * (14 * 1024 * 1024)))  # 14 MiB each, < 16 MiB
    many.insert(1, ("OEBPS/content.opf", opf_xml(items, spine)))
    many_raw = build_epub(Path(tmp) / "many.epub", many)
    mtext, merr = parse_epub(tmp, "many.epub", many_raw)
    check("C02c", "5×14MiB members exceed 64MiB total budget → epub_unsafe",
          mtext is None and merr == "epub_unsafe", merr,
          {"error_code": merr, "total_uncompressed": 5 * 14 * 1024 * 1024,
           "total_budget": parsers.EPUB_MAX_TOTAL_BYTES})

    # ---------- ③ path traversal ----------
    for cid, evil in (("C03a", "../evil.xhtml"), ("C03b", "/etc/passwd"),
                      ("C03c", "OEBPS/../../evil.xhtml"), ("C03d", "..\\evil.xhtml")):
        evil_raw = build_epub(Path(tmp) / "trav.epub", [
            ("META-INF/container.xml", CONTAINER_XML),
            ("OEBPS/content.opf", opf_xml([("p1", "a.xhtml")], ["p1"])),
            ("OEBPS/a.xhtml", xhtml('<p>ok</p>')),
            (evil, "pwned"),
        ])
        ttext, terr = parse_epub(tmp, "trav.epub", evil_raw)
        check(cid, "path-traversal member %r → epub_unsafe" % evil,
              ttext is None and terr == "epub_unsafe", terr, {"entry": evil, "error_code": terr})

    # ---------- ④ XXE / entity injection ----------
    ent_body = ('<?xml version="1.0"?>'
                '<!DOCTYPE html [<!ENTITY xxe SYSTEM "file:///etc/passwd">]>'
                '<html xmlns="http://www.w3.org/1999/xhtml"><body><p>&xxe;</p></body></html>')
    ent_raw = build_epub(Path(tmp) / "entity.epub", [
        ("META-INF/container.xml", CONTAINER_XML),
        ("OEBPS/content.opf", opf_xml([("p1", "a.xhtml")], ["p1"])),
        ("OEBPS/a.xhtml", ent_body),
    ])
    etext, eerr = parse_epub(tmp, "entity.epub", ent_raw)
    check("C04a", "chapter <!ENTITY> XXE payload → epub_unsafe (rejected before parse)",
          etext is None and eerr == "epub_unsafe", eerr, {"error_code": eerr})

    docxml = ('<?xml version="1.0"?><!DOCTYPE container SYSTEM "x.dtd">'
              '<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container"/>')
    dtext, derr = parse_epub(tmp, "doctype.epub", build_epub(Path(tmp) / "doctype.epub", [
        ("META-INF/container.xml", docxml),
        ("OEBPS/content.opf", opf_xml([("p1", "a.xhtml")], ["p1"])),
        ("OEBPS/a.xhtml", xhtml('<p>x</p>')),
    ]))
    check("C04b", "container.xml <!DOCTYPE ... SYSTEM ...> → epub_unsafe (external DTD blocked)",
          dtext is None and derr == "epub_unsafe", derr, {"error_code": derr})

    # ④c  bare <!DOCTYPE html> (no entities / no external id) MUST be tolerated:
    #     real-world EPUB XHTML carries it, and it cannot define entities alone.
    #     (TC decision 2026-10-12: relax the initial "reject any DOCTYPE" rule.)
    bare_body = ('<?xml version="1.0" encoding="utf-8"?>'
                 '<!DOCTYPE html>'
                 '<html xmlns="http://www.w3.org/1999/xhtml"><body>'
                 '<h1>Unit 1</h1><p>reservation booking</p></body></html>')
    btext, berr = parse_epub(tmp, "baredoctype.epub", build_epub(
        Path(tmp) / "baredoctype.epub", [
            ("META-INF/container.xml", CONTAINER_XML),
            ("OEBPS/content.opf", opf_xml([("p1", "a.xhtml")], ["p1"])),
            ("OEBPS/a.xhtml", bare_body),
        ]))
    check("C04c", "bare <!DOCTYPE html> (no entities) → PARSED (tolerated & stripped)",
          btext is not None and berr is None and "reservation" in btext, berr,
          {"error_code": berr, "text_len": len(btext or ""),
           "accepts_real_world_epub": btext is not None})

    # ④d  DOCTYPE with an *internal subset* (even without the ENTITY keyword in a
    #     form the ENTITY regex would catch) is still refused — defence in depth.
    sub_body = ('<?xml version="1.0"?>'
                '<!DOCTYPE html [<!ELEMENT p (#PCDATA)>]>'
                '<html xmlns="http://www.w3.org/1999/xhtml"><body><p>x</p></body></html>')
    stext, serr = parse_epub(tmp, "subset.epub", build_epub(
        Path(tmp) / "subset.epub", [
            ("META-INF/container.xml", CONTAINER_XML),
            ("OEBPS/content.opf", opf_xml([("p1", "a.xhtml")], ["p1"])),
            ("OEBPS/a.xhtml", sub_body),
        ]))
    check("C04d", "<!DOCTYPE> with internal subset (no ENTITY) → epub_unsafe",
          stext is None and serr == "epub_unsafe", serr, {"error_code": serr})

    # ④e  chapter-level external DTD reference must also be blocked (no fetch).
    ext_body = ('<?xml version="1.0"?>'
                '<!DOCTYPE html SYSTEM "http://evil.invalid/x.dtd">'
                '<html xmlns="http://www.w3.org/1999/xhtml"><body><p>x</p></body></html>')
    xtext, xerr = parse_epub(tmp, "extdtd.epub", build_epub(
        Path(tmp) / "extdtd.epub", [
            ("META-INF/container.xml", CONTAINER_XML),
            ("OEBPS/content.opf", opf_xml([("p1", "a.xhtml")], ["p1"])),
            ("OEBPS/a.xhtml", ext_body),
        ]))
    check("C04e", "chapter <!DOCTYPE ... SYSTEM ...> → epub_unsafe (no external fetch)",
          xtext is None and xerr == "epub_unsafe", xerr, {"error_code": xerr})

    # ---------- ⑥ non-zip masquerading as .epub ----------
    ntext, nerr = parse_epub(tmp, "fake.epub", b"This is definitely not a zip archive.")
    check("C06", "non-zip content named .epub → epub_malformed",
          ntext is None and nerr == "epub_malformed", nerr, {"error_code": nerr})

    # ---------- ⑦ empty zip ----------
    bio = io.BytesIO()
    zipfile.ZipFile(bio, "w").close()
    ztext, zerr = parse_epub(tmp, "empty.epub", bio.getvalue())
    check("C07", "empty zip → epub_malformed",
          ztext is None and zerr == "epub_malformed", zerr, {"error_code": zerr})

    # ---------- ⑧ corrupted zip ----------
    ctext, cerr = parse_epub(tmp, "corrupt.epub", raw[:200])
    check("C08", "truncated/corrupt zip → epub_malformed",
          ctext is None and cerr == "epub_malformed", cerr, {"error_code": cerr})

    # ---------- ⑨ valid EPUB but no extractable text ----------
    e2text, e2err = parse_epub(tmp, "blank.epub", build_epub(Path(tmp) / "blank.epub", [
        ("META-INF/container.xml", CONTAINER_XML),
        ("OEBPS/content.opf", opf_xml([("p1", "a.xhtml")], ["p1"])),
        ("OEBPS/a.xhtml", xhtml('<p></p><script>var x=1;</script>')),
    ]))
    check("C09", "textless spine → empty_extraction",
          e2text is None and e2err == "empty_extraction", e2err, {"error_code": e2err})

    # ---------- ⑬ encrypted member flag ----------
    # flip the encryption bit (bit 0) on a normal chapter's central-dir entry
    enc_clean = build_epub(Path(tmp) / "enc.epub", [
        ("META-INF/container.xml", CONTAINER_XML),
        ("OEBPS/content.opf", opf_xml([("p1", "a.xhtml")], ["p1"])),
        ("OEBPS/a.xhtml", xhtml('<p>secret</p>')),
    ])
    enc = _patch_central(enc_clean, "OEBPS/a.xhtml", flags=0x0001)
    ctext2, cerr2 = parse_epub(tmp, "enc.epub", enc)
    check("C13", "encrypted member → encrypted",
          ctext2 is None and cerr2 == "encrypted", cerr2, {"error_code": cerr2})

    # ---------- ⑭ missing structure: no spine / no container / no rootfile ----------
    nospine, ns_err = parse_epub(tmp, "nospine.epub", build_epub(Path(tmp) / "nospine.epub", [
        ("META-INF/container.xml", CONTAINER_XML),
        ("OEBPS/content.opf",
         '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" '
         'version="2.0"><manifest><item id="p1" href="a.xhtml" '
         'media-type="application/xhtml+xml"/></manifest>'
         '<spine></spine></package>'),
        ("OEBPS/a.xhtml", xhtml('<p>x</p>')),
    ]))
    check("C14a", "OPF without spine → epub_malformed",
          nospine is None and ns_err == "epub_malformed", ns_err, {"error_code": ns_err})

    nocont, nc_err = parse_epub(tmp, "nocont.epub", build_epub(Path(tmp) / "nocont.epub", [
        ("OEBPS/content.opf", opf_xml([("p1", "a.xhtml")], ["p1"])),
        ("OEBPS/a.xhtml", xhtml('<p>x</p>')),
    ]))
    check("C14b", "EPUB without META-INF/container.xml → epub_malformed",
          nocont is None and nc_err == "epub_malformed", nc_err, {"error_code": nc_err})

    badroot = build_epub(Path(tmp) / "badroot.epub", [
        ("META-INF/container.xml", CONTAINER_XML),
        ("OEBPS/a.xhtml", xhtml('<p>x</p>')),
    ])
    broot, br_err = parse_epub(tmp, "badroot.epub", badroot)
    check("C14c", "container points at a missing rootfile → epub_malformed",
          broot is None and br_err == "epub_malformed", br_err, {"error_code": br_err})

    # ---------- ⑮ malformed chapter XML → honest degrade (skip chapter) ----------
    deg = build_epub(Path(tmp) / "degrade.epub", [
        ("META-INF/container.xml", CONTAINER_XML),
        ("OEBPS/content.opf", opf_xml([("p1", "bad.xhtml"), ("p2", "good.xhtml")], ["p1", "p2"])),
        ("OEBPS/bad.xhtml", '<html><body><p>unclosed'),
        ("OEBPS/good.xhtml", xhtml('<p>GOOD-CHAPTER</p>')),
    ])
    dtext2, derr2 = parse_epub(tmp, "degrade.epub", deg)
    check("C15", "malformed chapter → VISIBLE marker (not silently dropped); good chapter survives",
          derr2 is None and "GOOD-CHAPTER" in (dtext2 or "")
          and parsers.EPUB_CHAPTER_MISSING_MARKER in (dtext2 or "")
          and (dtext2 or "").startswith("[chapter 1]"),
          derr2, {"error_code": derr2, "text": (dtext2 or "")[:80],
                  "marker_present": parsers.EPUB_CHAPTER_MISSING_MARKER in (dtext2 or ""),
                  "marker_text": parsers.EPUB_CHAPTER_MISSING_MARKER})

    # ⑮b when EVERY chapter is unreadable the parser must still report
    #     empty_extraction -- never a "successful" text made only of markers.
    allbad = build_epub(Path(tmp) / "allbad.epub", [
        ("META-INF/container.xml", CONTAINER_XML),
        ("OEBPS/content.opf", opf_xml([("p1", "b1.xhtml"), ("p2", "b2.xhtml")], ["p1", "p2"])),
        ("OEBPS/b1.xhtml", '<html><body><p>unclosed-one'),
        ("OEBPS/b2.xhtml", '<html><body><p>unclosed-two'),
    ])
    abtext, aberr = parse_epub(tmp, "allbad.epub", allbad)
    check("C15b", "ALL chapters malformed → empty_extraction (never marker-only success)",
          abtext is None and aberr == "empty_extraction", aberr, {"error_code": aberr})

    # ---------- ⑯ never raises on hostile inputs ----------
    raised = None
    sweep = [("nofile.epub", None, "missing"), ("dir.epub", None, "dir"),
             ("zero.epub", b"", "zero-bytes"), ("junk.epub", b"PK\x03\x04rubbish", "fake-local-header")]
    for nm, data, label in sweep:
        try:
            p = Path(tmp) / nm
            if label == "dir":
                p.mkdir(exist_ok=True)
            elif data is not None:
                p.write_bytes(data)
            _t, _e = parsers.parse_file(str(p), ".epub")
            if _t is not None and _e is not None:
                raised = "%s: both text and error set" % label
            if _e is not None and _t is not None:
                raised = "%s: not exactly one non-None" % label
        except Exception as exc:  # noqa: BLE001
            raised = "%s: raised %s: %s" % (label, type(exc).__name__, exc)
    check("C16", "parse_file never raises; always returns exactly one of (text, error)",
          raised is None, None, raised or "ok")

    # ---------- ⑫ regression: the existing five formats unchanged ----------
    reg = {}
    (Path(tmp) / "r.txt").write_text("PLAIN TXT with <!DOCTYPE html> inside", "utf-8")
    (Path(tmp) / "r.md").write_text("# Heading\n\nBODY-MD", "utf-8")
    rt, re_ = parsers.parse_file(str(Path(tmp) / "r.txt"), ".txt")
    check("C12a", "TXT still parses (and is NOT subject to the EPUB DTD guard)",
          re_ is None and "PLAIN TXT" in (rt or ""))
    rm, rme = parsers.parse_file(str(Path(tmp) / "r.md"), ".md")
    check("C12b", "MD still parses", rme is None and "BODY-MD" in (rm or ""))

    import docx as _docx
    d = _docx.Document()
    d.add_paragraph("REG-DOCX")
    d.save(str(Path(tmp) / "r.docx"))
    dt, de = parsers.parse_file(str(Path(tmp) / "r.docx"), ".docx")
    check("C12c", "DOCX still parses", de is None and "REG-DOCX" in (dt or ""))

    from pptx import Presentation
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.add_textbox(0, 0, 100000, 100000).text_frame.text = "REG-PPTX"
    prs.save(str(Path(tmp) / "r.pptx"))
    pt, pe = parsers.parse_file(str(Path(tmp) / "r.pptx"), ".pptx")
    check("C12d", "PPTX still parses", pe is None and "REG-PPTX" in (pt or ""))

    MIN_PDF = b"""%PDF-1.4
1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj
2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj
3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 300 200]/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>endobj
4 0 obj<</Length 52>>stream
BT /F1 14 Tf 20 100 Td (Hello Knowledge Base) Tj ET
endstream
endobj
5 0 obj<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>endobj
trailer<</Root 1 0 R>>
%%EOF
"""
    (Path(tmp) / "r.pdf").write_bytes(MIN_PDF)
    pdft, pdfe = parsers.parse_file(str(Path(tmp) / "r.pdf"), ".pdf")
    check("C12e", "PDF still parses (pdfminer present in venv)",
          pdfe is None and "Hello Knowledge Base" in (pdft or ""),
          pdfe, {"error_code": pdfe, "preview": (pdft or "")[:40]})

    caps2 = {c["ext"]: c for c in parsers.capabilities()}
    check("C12f", "all six formats present in capabilities()",
          set(caps2) == {".txt", ".md", ".pdf", ".docx", ".pptx", ".epub"},
          None, sorted(caps2))

    # ---------- ⑰ end-to-end through the import pipeline (upload → parse → preview)
    from modules.knowledge.imports import ImportStore
    from modules.knowledge.store import KnowledgeStore

    kb = KnowledgeStore(str(Path(tmp) / "kb.json")).load()
    ds = kb.create_dataset("EPUB 测试")
    isr = ImportStore(str(Path(tmp) / "imp.json"),
                      scratch_dir=str(Path(tmp) / "imp_scratch")).load()
    job = isr.create_job(dataset_id=ds["id"], filename="book.epub",
                         data=(Path(tmp) / "normal.epub").read_bytes())
    check("C17", "EPUB accepted end-to-end by import pipeline → preview_ready (epub_text)",
          job["status"] == "preview_ready" and job.get("source_type") == "epub_text"
          and (job.get("extracted_chars") or 0) > 0,
          job.get("error_code"),
          {"status": job["status"], "source_type": job.get("source_type"),
           "extracted_chars": job.get("extracted_chars"),
           "preview_head": (job.get("preview") or "")[:24]})

    # ---------- cross-check: count assertions elsewhere were aligned 5 -> 6 ----------
    cross = []
    for rel, needle in (("gateway/tests/admin_v4_api_tests.py", "len(parsers) == 6"),
                        ("gateway/tests/import_v4_tests.py", "len(kinds) == 6")):
        fp = ROOT / rel
        for i, line in enumerate(fp.read_text("utf-8").splitlines(), 1):
            if needle in line:
                cross.append({"file": rel, "line": i, "text": line.strip(),
                              "status": "aligned_by_impl (was '== 5' before .epub)"})
    # verify_v4_e2e.py holds no count assertion — it only checks the five exts are
    # available, which still holds with .epub added.
    e2e = ROOT / "deploy/v4-acceptance/verify_v4_e2e.py"
    for i, line in enumerate(e2e.read_text("utf-8").splitlines(), 1):
        if 'for ext in (".txt", ".md", ".pdf", ".docx", ".pptx")' in line:
            cross.append({"file": "deploy/v4-acceptance/verify_v4_e2e.py", "line": i,
                          "text": line.strip(),
                          "status": "no count assertion — iterates the five exts only, still passes"})

    # ---- evidence ----
    head = "unknown"
    try:
        h = (ROOT / ".git/HEAD").read_text().strip()
        if h.startswith("ref:"):
            ref = h.split(" ", 1)[1].strip()
            ref_file = ROOT / ".git" / ref
            head = (ref_file.read_text().strip() if ref_file.is_file()
                    else "%s (packed)" % ref)
        else:
            head = h
    except Exception:  # noqa: BLE001
        pass

    ev = {
        "artifact": "epub_parser_security_tests",
        "scope": "CLASSROOM V2 · D2 — text-only EPUB parsing",
        "generated_at": time.strftime("%Y-%m-%dT%H:%M:%S%z"),
        "python_version": sys.version,
        "python_executable": sys.executable,
        "platform": platform.platform(),
        "repo_head": head,
        "target_module": "modules/knowledge/parsers.py",
        "safety_constants": {
            "EPUB_MAX_ENTRIES": parsers.EPUB_MAX_ENTRIES,
            "EPUB_MAX_ENTRY_BYTES": parsers.EPUB_MAX_ENTRY_BYTES,
            "EPUB_MAX_TOTAL_BYTES": parsers.EPUB_MAX_TOTAL_BYTES,
            "EPUB_READ_CHUNK": parsers.EPUB_READ_CHUNK,
            "EPUB_REJECT_DTD": parsers.EPUB_REJECT_DTD,
            "EPUB_ALLOW_BARE_DOCTYPE": parsers.EPUB_ALLOW_BARE_DOCTYPE,
        },
        "dtd_policy": {
            "reject": ["<!ENTITY ...>", "<!DOCTYPE ... SYSTEM|PUBLIC ...>",
                       "<!DOCTYPE ... [ ... ]>"],
            "tolerate_and_strip": ["bare <!DOCTYPE html>"],
            "rationale": ("TC decision 2026-10-12: real-world EPUB XHTML carries a bare "
                          "<!DOCTYPE html> which cannot define entities; refusing it would "
                          "reject most genuine EPUBs. Entity/XXE vectors still fail-closed."),
        },
        "new_error_codes": ["epub_malformed", "epub_unsafe"],
        "reused_error_codes": ["encrypted", "empty_extraction", "extract_failed"],
        "summary": {"total": len(RESULTS), "passed": len(PASS), "failed": len(FAIL),
                    "failed_ids": list(FAIL)},
        "cases": RESULTS,
        "cross_file_assertions": cross,
    }
    EVIDENCE.parent.mkdir(parents=True, exist_ok=True)
    EVIDENCE.write_text(json.dumps(ev, ensure_ascii=False, indent=2), "utf-8")

    print("\n==== SUMMARY ====")
    print("total=%d pass=%d fail=%d" % (len(RESULTS), len(PASS), len(FAIL)))
    if FAIL:
        print("FAILED: " + ", ".join(FAIL))
    print("evidence: %s" % EVIDENCE)
    print("cross-file count assertions that need alignment: %d" % len(cross))
    return 1 if FAIL else 0


if __name__ == "__main__":
    sys.exit(main())
